from __future__ import annotations

import io
import re
from typing import Any, Iterable

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BRAND_NAVY = RGBColor(0x0B, 0x1F, 0x3A)
BRAND_BLUE = RGBColor(0x1E, 0x88, 0xE5)
TEXT_DARK = RGBColor(0x21, 0x29, 0x2F)
TEXT_MUTED = RGBColor(0x5F, 0x6B, 0x76)
REDACTION_RED = RGBColor(0xB7, 0x1C, 0x1C)
LIGHT_BLUE_HEX = "EAF3FB"
LIGHT_GREY_HEX = "F5F7F9"
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _set_run_font(run, name: str = "Calibri") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:shd"))
    if existing is not None:
        tc_pr.remove(existing)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def _set_cell_margins(cell, top: int = 100, start: int = 120,
                      bottom: int = 100, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for margin, value in (
        ("top", top), ("start", start), ("bottom", bottom), ("end", end)
    ):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])
    _set_run_font(run)
    run.font.size = Pt(8)
    run.font.color.rgb = TEXT_MUTED


def _new_document(title: str, subject: str, transcript_id: str) -> Document:
    doc = Document()
    doc.core_properties.title = title
    doc.core_properties.subject = subject
    doc.core_properties.author = "Stelic"

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = TEXT_DARK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")

    heading_1 = doc.styles["Heading 1"]
    heading_1.font.name = "Calibri"
    heading_1.font.size = Pt(16)
    heading_1.font.bold = True
    heading_1.font.color.rgb = BRAND_BLUE
    heading_1.paragraph_format.space_before = Pt(16)
    heading_1.paragraph_format.space_after = Pt(8)
    heading_1.paragraph_format.keep_with_next = True
    heading_1._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    heading_1._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")

    heading_2 = doc.styles["Heading 2"]
    heading_2.font.name = "Calibri"
    heading_2.font.size = Pt(13)
    heading_2.font.bold = True
    heading_2.font.color.rgb = BRAND_BLUE
    heading_2.paragraph_format.space_before = Pt(12)
    heading_2.paragraph_format.space_after = Pt(6)
    heading_2.paragraph_format.keep_with_next = True
    heading_2._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    heading_2._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")

    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header_run = header.add_run("STELIC  |  RECRUITING")
        _set_run_font(header_run)
        header_run.bold = True
        header_run.font.size = Pt(8)
        header_run.font.color.rgb = BRAND_NAVY

        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer_run = footer.add_run(
            f"Confidential  •  Fireflies transcript {transcript_id}  •  Page "
        )
        _set_run_font(footer_run)
        footer_run.font.size = Pt(8)
        footer_run.font.color.rgb = TEXT_MUTED
        _add_page_number(footer)

    return doc


def _add_title_banner(doc: Document, title: str, subtitle: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), "0B1F3A")
    p_pr.append(shading)
    run = paragraph.add_run(title)
    _set_run_font(run, "Calibri")
    run.bold = True
    run.font.size = Pt(23)
    run.font.color.rgb = WHITE

    if subtitle:
        subtitle_paragraph = doc.add_paragraph()
        subtitle_paragraph.paragraph_format.space_before = Pt(0)
        subtitle_paragraph.paragraph_format.space_after = Pt(14)
        subtitle_run = subtitle_paragraph.add_run(subtitle)
        _set_run_font(subtitle_run, "Calibri")
        subtitle_run.bold = True
        subtitle_run.font.size = Pt(14)
        subtitle_run.font.color.rgb = BRAND_NAVY


def _set_table_geometry(table, column_widths: list[int]) -> None:
    total_width = sum(column_widths)
    tbl_pr = table._tbl.tblPr

    width = tbl_pr.first_child_found_in("w:tblW")
    if width is None:
        width = OxmlElement("w:tblW")
        tbl_pr.append(width)
    width.set(qn("w:w"), str(total_width))
    width.set(qn("w:type"), "dxa")

    indent = tbl_pr.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        tbl_pr.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid_columns = list(table._tbl.tblGrid)
    for grid_column, column_width in zip(grid_columns, column_widths):
        grid_column.set(qn("w:w"), str(column_width))

    for row in table.rows:
        for cell, column_width in zip(row.cells, column_widths):
            tc_width = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_width is None:
                tc_width = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_width)
            tc_width.set(qn("w:w"), str(column_width))
            tc_width.set(qn("w:type"), "dxa")


def _add_metadata_table(doc: Document, rows: Iterable[tuple[str, Any]]) -> None:
    clean_rows = [(label, str(value).strip()) for label, value in rows if value]
    if not clean_rows:
        return

    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False

    for index, (label, value) in enumerate(clean_rows):
        cells = table.add_row().cells
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _shade_cell(cells[0], LIGHT_BLUE_HEX)
        if index % 2:
            _shade_cell(cells[1], LIGHT_GREY_HEX)
        _set_cell_margins(cells[0])
        _set_cell_margins(cells[1])

        label_run = cells[0].paragraphs[0].add_run(label)
        _set_run_font(label_run)
        label_run.bold = True
        label_run.font.size = Pt(9)
        label_run.font.color.rgb = BRAND_NAVY

        value_run = cells[1].paragraphs[0].add_run(value)
        _set_run_font(value_run)
        value_run.font.size = Pt(9)
        value_run.font.color.rgb = TEXT_DARK

    _set_table_geometry(table, [2700, 6660])
    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def _add_section_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Heading 1")
    run = paragraph.add_run(text)
    _set_run_font(run, "Calibri")


def _external_email(emails: list[str]) -> str:
    normalized = [str(email).strip() for email in emails if str(email).strip()]
    return next(
        (email for email in normalized if not email.lower().endswith("@stelic.com")),
        normalized[0] if normalized else "",
    )


def _candidate_from_title(title: str) -> str:
    value = str(title or "").strip()
    for prefix in ("Interview Transcript -", "Interview Summary -"):
        if value.lower().startswith(prefix.lower()):
            return value[len(prefix):].strip()
    return value


def _duration_label(value: Any) -> str:
    try:
        minutes = int(value)
    except (TypeError, ValueError):
        return ""
    return f"{minutes} minute" if minutes == 1 else f"{minutes} minutes"


def _parse_transcript(raw_text: str) -> tuple[dict[str, str], str]:
    normalized = str(raw_text or "").replace("\r\n", "\n").replace("\r", "\n")
    metadata: dict[str, str] = {}
    body = normalized.strip()

    divider = re.search(r"(?m)^-{8,}\s*$", normalized)
    if divider:
        header = normalized[:divider.start()]
        body = normalized[divider.end():].strip()
    else:
        header = ""

    for line in header.splitlines():
        if ":" not in line:
            continue
        key, value = line.split(":", 1)
        key = key.strip().lower()
        if key in {"candidate", "email", "meeting", "date", "fireflies"}:
            metadata[key] = value.strip()

    return metadata, body


def _transcript_turns(body: str) -> list[tuple[str, str]]:
    pattern = re.compile(r"(?m)^(?P<speaker>[^\n:]{1,100}):\s*\n")
    matches = list(pattern.finditer(body))
    if not matches:
        return [("Transcript", body.strip())] if body.strip() else []

    turns: list[tuple[str, str]] = []
    for index, match in enumerate(matches):
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(body)
        text = body[start:end].strip()
        turns.append((match.group("speaker").strip(), text))
    return turns


def _add_transcript_turn(doc: Document, speaker: str, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.left_indent = Inches(0.12)

    speaker_run = paragraph.add_run(f"{speaker}: ")
    _set_run_font(speaker_run)
    speaker_run.bold = True
    speaker_run.font.color.rgb = BRAND_BLUE

    text_run = paragraph.add_run(text or "—")
    _set_run_font(text_run)
    text_run.font.color.rgb = TEXT_DARK
    if "[REDACTED" in (text or "").upper():
        text_run.italic = True
        text_run.font.color.rgb = REDACTION_RED


def _clean_list_line(value: str) -> str:
    return re.sub(r"^\s*(?:[-–—•*]|\d+[.)])\s*", "", value).strip()


def _add_body_paragraphs(doc: Document, text: str) -> None:
    blocks = [block.strip() for block in re.split(r"\n\s*\n", text or "") if block.strip()]
    if not blocks:
        doc.add_paragraph("No content was returned by Fireflies.").italic = True
        return
    for block in blocks:
        paragraph = doc.add_paragraph(block)
        paragraph.paragraph_format.space_after = Pt(7)


def _add_bullets(doc: Document, text: str) -> None:
    items = [_clean_list_line(line) for line in (text or "").splitlines()]
    items = [item for item in items if item]
    if not items:
        doc.add_paragraph("None provided.").italic = True
        return
    for item in items:
        paragraph = doc.add_paragraph(item, style="List Bullet")
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.line_spacing = 1.167


def _save_document(doc: Document) -> bytes:
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def build_interview_transcript_docx(payload: dict[str, Any]) -> bytes:
    meeting = payload.get("meeting") or {}
    notes = payload.get("meeting_notes") or []
    raw_text = "\n\n".join(
        str(note.get("discussion") or "")
        for note in notes
        if isinstance(note, dict)
    ).strip()
    metadata, transcript_body = _parse_transcript(raw_text)

    candidate = metadata.get("candidate") or _candidate_from_title(meeting.get("title", ""))
    candidate_email = metadata.get("email") or _external_email(
        meeting.get("participant_emails") or []
    )
    transcript_id = str(meeting.get("transcript_id") or "").strip()

    doc = _new_document(
        title=f"Interview Transcript - {candidate}",
        subject="Stelic interview transcript",
        transcript_id=transcript_id,
    )
    _add_title_banner(doc, "INTERVIEW TRANSCRIPT", candidate)
    _add_metadata_table(doc, [
        ("Candidate", candidate),
        ("Email", candidate_email),
        ("Meeting", metadata.get("meeting")),
        ("Date", metadata.get("date") or meeting.get("date")),
        ("Duration", _duration_label(meeting.get("duration_minutes"))),
        ("Attendees", ", ".join(meeting.get("attendees") or [])),
        ("Fireflies", metadata.get("fireflies")),
        ("Transcript ID", transcript_id),
    ])

    _add_section_heading(doc, "Transcript")
    turns = _transcript_turns(transcript_body)
    if turns:
        for speaker, text in turns:
            _add_transcript_turn(doc, speaker, text)
    else:
        doc.add_paragraph("No transcript content was returned by Fireflies.").italic = True

    return _save_document(doc)


def build_interview_summary_docx(payload: dict[str, Any]) -> bytes:
    meeting = payload.get("meeting") or {}
    notes = payload.get("meeting_notes") or []
    candidate = _candidate_from_title(meeting.get("title", ""))
    candidate_email = _external_email(meeting.get("participant_emails") or [])
    transcript_id = str(meeting.get("transcript_id") or "").strip()

    doc = _new_document(
        title=f"Interview Summary - {candidate}",
        subject="Stelic interview summary",
        transcript_id=transcript_id,
    )
    _add_title_banner(doc, "INTERVIEW SUMMARY", candidate)
    _add_metadata_table(doc, [
        ("Candidate", candidate),
        ("Email", candidate_email),
        ("Date", meeting.get("date")),
        ("Duration", _duration_label(meeting.get("duration_minutes"))),
        ("Attendees", ", ".join(meeting.get("attendees") or [])),
        ("Transcript ID", transcript_id),
    ])

    rendered = False
    for note in notes:
        if not isinstance(note, dict):
            continue
        topic = str(note.get("topic") or "Summary").strip()
        discussion = str(note.get("discussion") or "").strip()
        if not discussion:
            continue
        rendered = True
        _add_section_heading(doc, topic)
        if topic.lower() in {"topics discussed", "action items", "actions", "next steps"}:
            _add_bullets(doc, discussion)
        else:
            _add_body_paragraphs(doc, discussion)

    if not rendered:
        _add_section_heading(doc, "Summary")
        doc.add_paragraph("No summary content was returned by Fireflies.").italic = True

    return _save_document(doc)
