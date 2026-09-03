import io
import unittest

from docx import Document

from interview_builder import (
    build_interview_summary_docx,
    build_interview_transcript_docx,
)
from main import InterviewArtifactRequest, app, safe_download_filename
from minutes_builder import build_minutes_docx


def document_text(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    values = [paragraph.text for paragraph in document.paragraphs]

    for table in document.tables:
        for row in table.rows:
            values.extend(cell.text for cell in row.cells)

    for section in document.sections:
        values.extend(paragraph.text for paragraph in section.header.paragraphs)
        values.extend(paragraph.text for paragraph in section.footer.paragraphs)

    return "\n".join(values)


def base_meeting(title: str) -> dict:
    return {
        "title": title,
        "date": "September 3, 2026 at 10:00 AM ET",
        "date_iso": "2026-09-03T14:00:00.000Z",
        "duration_minutes": 45,
        "attendees": ["Rob Crouch", "Jane Doe"],
        "participant_emails": [
            "rcrouch@stelic.com",
            "jane@example.com",
        ],
        "transcript_id": "01TESTTRANSCRIPT",
        "meeting_type": "other",
        "type_confidence": None,
        "type_reasoning": "Fireflies interview artifact",
    }


class InterviewBuilderTests(unittest.TestCase):
    def test_transcript_has_interview_layout_and_clean_body(self):
        payload = {
            "meeting": base_meeting("Interview Transcript - Jane Doe"),
            "meeting_notes": [{
                "topic": "Full Transcript",
                "discussion": (
                    "STELIC INTERVIEW TRANSCRIPT\n\n"
                    "Candidate: Jane Doe\n"
                    "Email: jane@example.com\n"
                    "Meeting: Project Manager Interview\n"
                    "Date: September 3, 2026 at 10:00 AM ET\n"
                    "Fireflies: https://app.fireflies.ai/view/test\n\n"
                    "----------------------------------------\n\n"
                    "Rob Crouch:\nWelcome to the interview.\n\n"
                    "Jane Doe:\nThank you. I am glad to be here.\n\n"
                    "Rob Crouch:\n[REDACTED — SECURITY CREDENTIALS]"
                ),
                "decision": "",
                "open_questions": [],
            }],
            "filename": "Interview Transcript - Jane Doe.docx",
        }

        content = build_interview_transcript_docx(payload)
        text = document_text(content)

        self.assertTrue(content.startswith(b"PK"))
        self.assertIn("INTERVIEW TRANSCRIPT", text)
        self.assertIn("Jane Doe", text)
        self.assertIn("Project Manager Interview", text)
        self.assertIn("Welcome to the interview.", text)
        self.assertIn("[REDACTED — SECURITY CREDENTIALS]", text)
        self.assertNotIn("MEETING NOTES", text)
        self.assertNotIn("Classified as other", text)

    def test_summary_has_interview_sections_and_bullets(self):
        payload = {
            "meeting": base_meeting("Interview Summary - Jane Doe"),
            "meeting_notes": [
                {
                    "topic": "Overview",
                    "discussion": "Jane has strong project controls experience.",
                    "decision": "",
                    "open_questions": [],
                },
                {
                    "topic": "Notes",
                    "discussion": "She has worked on two major programs.",
                    "decision": "",
                    "open_questions": [],
                },
                {
                    "topic": "Topics Discussed",
                    "discussion": "Scheduling\nCost control",
                    "decision": "",
                    "open_questions": [],
                },
                {
                    "topic": "Action Items",
                    "discussion": "Rob to follow up\nSchedule technical review",
                    "decision": "",
                    "open_questions": [],
                },
            ],
            "filename": "Interview Summary - Jane Doe.docx",
        }

        content = build_interview_summary_docx(payload)
        text = document_text(content)

        self.assertTrue(content.startswith(b"PK"))
        self.assertIn("INTERVIEW SUMMARY", text)
        self.assertIn("Overview", text)
        self.assertIn("Topics Discussed", text)
        self.assertIn("Action Items", text)
        self.assertIn("Schedule technical review", text)
        self.assertNotIn("MEETING NOTES", text)
        self.assertNotIn("Classified as other", text)

    def test_current_n8n_payload_allows_extra_fields(self):
        payload = {
            "meeting": base_meeting("Interview Summary - Jane Doe"),
            "meeting_notes": [],
            "filename": "summary.docx",
            "scoring": None,
            "todos": None,
            "retreat_items": [],
            "agenda": None,
        }

        request = InterviewArtifactRequest(**payload)
        self.assertEqual(request.filename, "summary.docx")

    def test_routes_are_registered(self):
        paths = {route.path for route in app.routes}
        self.assertIn("/build-minutes", paths)
        self.assertIn("/build-interview-transcript", paths)
        self.assertIn("/build-interview-summary", paths)
        self.assertIn("/health", paths)

    def test_filename_is_sanitized_and_forced_to_docx(self):
        self.assertEqual(
            safe_download_filename('../bad"name.txt', "fallback.docx"),
            "bad-name.txt.docx",
        )

    def test_existing_minutes_builder_still_generates_docx(self):
        payload = {
            "meeting": base_meeting("Weekly Operations Meeting"),
            "meeting_notes": [],
            "scoring": None,
            "todos": None,
            "retreat_items": [],
            "agenda": None,
            "filename": "minutes.docx",
        }

        content = build_minutes_docx(payload)
        self.assertTrue(content.startswith(b"PK"))
        self.assertIn("Meeting Notes", document_text(content))


if __name__ == "__main__":
    unittest.main()
