from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import io


# Stelic brand-adjacent colors (adjust to match if you have brand hex)
BRAND_DARK = RGBColor(0x0B, 0x1F, 0x3A)   # dark navy
BRAND_ACCENT = RGBColor(0x1E, 0x88, 0xE5)  # blue
GREEN = RGBColor(0x2E, 0x7D, 0x32)
AMBER = RGBColor(0xF9, 0xA8, 0x25)
RED = RGBColor(0xC6, 0x28, 0x28)
GREY = RGBColor(0x61, 0x61, 0x61)


def _shade_cell(cell, hex_color: str):
    """Apply background color to a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)


def _score_color(score: float) -> str:
    if score >= 8:
        return "C8E6C9"  # green tint
    if score >= 5:
        return "FFF9C4"  # yellow tint
    if score > 0:
        return "FFCDD2"  # red tint
    return "F5F5F5"      # grey for zero/no evidence


def _add_heading(doc: Document, text: str, level: int = 1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = BRAND_DARK
    return h


def _add_kv_row(table, key: str, value: str):
    row = table.add_row()
    row.cells[0].text = key
    row.cells[1].text = value
    # Set explicit widths so long values wrap instead of overflowing
    row.cells[0].width = Cm(3.5)
    row.cells[1].width = Cm(12.5)
    # Bold the key
    for para in row.cells[0].paragraphs:
        for run in para.runs:
            run.bold = True


# Rubric labels per meeting type — key mappings for pretty display
RUBRIC_LABELS = {
    "weekly_l10": {
        "on_time": "On-time start and end",
        "segue": "Segue",
        "scorecard": "Scorecard review",
        "rock_review": "Rock review",
        "headlines": "Headlines",
        "todo_review": "To-Do review",
        "ids_discipline": "IDS discipline",
        "cascading_messages": "Cascading messages",
        "conclude_and_rate": "Conclude and rate"
    },
    "quarterly_pulsing": {
        "segue": "Segue",
        "rock_completion": "Prior-quarter Rock completion",
        "ids_discipline": "IDS discipline",
        "new_rocks": "Next-quarter Rocks set",
        "vto_review": "V/TO review",
        "conclude": "Conclude and next steps"
    },
    "annual_planning": {
        "vto_refresh": "V/TO refresh",
        "prior_year_review": "Prior-year review",
        "three_year_picture": "3-Year Picture",
        "one_year_plan": "1-Year Plan",
        "q1_rocks": "Q1 Rocks",
        "ids_discipline": "IDS discipline",
        "team_health": "Team health"
    },
    "vision_building": {
        "core_values": "Core Values",
        "core_focus": "Core Focus",
        "ten_year_target": "10-Year Target",
        "marketing_strategy": "Marketing Strategy",
        "three_year_picture": "3-Year Picture"
    },
    "same_page": {
        "gwc": "GWC check-in",
        "personal_business_check_in": "Personal + business check-in",
        "issues_ids": "Business issues (IDS)",
        "commitments": "Commitments for next meeting"
    }
}

TYPE_TITLE_LABELS = {
    "weekly_l10": "Weekly L10 Minutes",
    "quarterly_pulsing": "Quarterly Pulsing Minutes",
    "annual_planning": "Annual Planning Minutes",
    "vision_building": "Vision Building Session Minutes",
    "same_page": "Same-Page Meeting Notes",
    "other": "Meeting Notes"
}


def build_minutes_docx(payload: dict) -> bytes:
    doc = Document()

    # Global styles
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    meeting = payload['meeting']
    meeting_type = meeting.get('meeting_type', 'other')
    scoring = payload.get('scoring')  # may be None for 'other'
    meeting_notes = payload.get('meeting_notes', [])
    todos = payload.get('todos')
    retreat_items = payload.get('retreat_items', [])
    agenda = payload.get('agenda') or {}

    # ---- TITLE BLOCK ----
    title = doc.add_paragraph()
    run = title.add_run(TYPE_TITLE_LABELS.get(meeting_type, "Meeting Notes"))
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = BRAND_DARK

    subtitle = doc.add_paragraph()
    sub_run = subtitle.add_run(meeting['title'])
    sub_run.font.size = Pt(14)
    sub_run.font.color.rgb = GREY

    # Metadata table
    meta_table = doc.add_table(rows=0, cols=2)
    meta_table.autofit = False
    meta_table.allow_autofit = False
    _add_kv_row(meta_table, "Date", meeting['date'])
    _add_kv_row(meta_table, "Duration", f"{meeting['duration_minutes']} minutes")
    _add_kv_row(meta_table, "Attendees", ", ".join(meeting['attendees']))
    _add_kv_row(meta_table, "Meeting Type", TYPE_TITLE_LABELS.get(meeting_type, meeting_type))
    if meeting.get('type_confidence') is not None:
        conf_pct = int(meeting['type_confidence'] * 100)
        _add_kv_row(meta_table, "Type Confidence", f"{conf_pct}%")

    doc.add_paragraph()

    # ---- SCORE (skipped for 'other') ----
    if scoring is not None and meeting_type in RUBRIC_LABELS:
        _add_heading(doc, f"{TYPE_TITLE_LABELS[meeting_type].replace(' Minutes', '')} Score", level=1)

        overall_p = doc.add_paragraph()
        overall_run = overall_p.add_run(f"Overall: {scoring['overall_score']} / 10")
        overall_run.font.size = Pt(18)
        overall_run.font.bold = True
        score_val = scoring['overall_score']
        if score_val >= 8:
            overall_run.font.color.rgb = GREEN
        elif score_val >= 5:
            overall_run.font.color.rgb = AMBER
        else:
            overall_run.font.color.rgb = RED

        summary_p = doc.add_paragraph()
        summary_run = summary_p.add_run(scoring['one_line_summary'])
        summary_run.italic = True
        summary_run.font.color.rgb = GREY

        # Subscore table
        score_table = doc.add_table(rows=1, cols=3)
        score_table.style = 'Light Grid Accent 1'
        hdr = score_table.rows[0].cells
        hdr[0].text = "Category"
        hdr[1].text = "Score"
        hdr[2].text = "Evidence"
        for cell in hdr:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True
                    run.font.color.rgb = BRAND_DARK

        labels = RUBRIC_LABELS[meeting_type]
        for key, label in labels.items():
            sub = scoring['subscores'].get(key, {"score": 0, "evidence": "no evidence"})
            row = score_table.add_row()
            row.cells[0].text = label
            row.cells[1].text = f"{sub['score']}"
            row.cells[2].text = sub.get('evidence', '')
            _shade_cell(row.cells[1], _score_color(sub['score']))

        if scoring.get('flags'):
            doc.add_paragraph()
            _add_heading(doc, "Flags", level=2)
            for flag in scoring['flags']:
                doc.add_paragraph(flag, style='List Bullet')

    # ---- MEETING NOTES ----
    if meeting_notes:
        doc.add_page_break()
        _add_heading(doc, "Meeting Notes", level=1)
        doc.add_paragraph(
            "Narrative summary of topics discussed, organized by subject."
        ).italic = True

        for note in meeting_notes:
            _add_heading(doc, note.get('topic', 'Untitled Topic'), level=2)

            if note.get('discussion'):
                doc.add_paragraph(note['discussion'])

            if note.get('decision'):
                p = doc.add_paragraph()
                label = p.add_run("Decision: ")
                label.bold = True
                label.font.color.rgb = BRAND_DARK
                p.add_run(note['decision'])

            open_qs = note.get('open_questions', [])
            if open_qs:
                p = doc.add_paragraph()
                label = p.add_run("Open questions:")
                label.bold = True
                label.font.color.rgb = BRAND_DARK
                for q in open_qs:
                    doc.add_paragraph(q, style='List Bullet')

    # ---- TO-DOS ----
    if todos is not None:
        doc.add_page_break()
        _add_heading(doc, "To-Dos", level=1)
        if todos.get('total_count', 0) == 0:
            doc.add_paragraph("No to-dos captured this meeting.").italic = True
        else:
            _render_todo_section(doc, "New To-Dos", todos.get('new', []))
            _render_todo_section(doc, "Carryover To-Dos", todos.get('carryover', []))

    # ---- RETREAT ITEMS ----
    if retreat_items:
        doc.add_page_break()
        _add_heading(doc, "Retreat Items Flagged This Meeting", level=1)
        doc.add_paragraph(
            "These items were flagged as strategic and appended to the running retreat list."
        ).italic = True

        for item in retreat_items:
            p = doc.add_paragraph(style='List Bullet')
            topic_run = p.add_run(item['topic'])
            topic_run.bold = True
            p.add_run(f" — {item['context']}")
            if item.get('raised_by'):
                p.add_run(f" (raised by {item['raised_by']})").italic = True

    # ---- NEXT AGENDA (only for recurring meeting types) ----
    if meeting_type in ('weekly_l10', 'quarterly_pulsing', 'same_page'):
        doc.add_page_break()
        _add_heading(doc, "Draft Agenda: Next Meeting", level=1)

        agenda_sections = [
            ("Scorecard items to revisit", agenda.get('carryover_scorecard_items', [])),
            ("Rocks off-track", agenda.get('carryover_rocks_offtrack', [])),
            ("Open to-dos for review", agenda.get('open_todos_for_review', [])),
            ("Unresolved issues", agenda.get('unresolved_issues', [])),
            ("Tabled items", agenda.get('tabled_items', []))
        ]

        for section_title, items in agenda_sections:
            _add_heading(doc, section_title, level=2)
            if not items:
                p = doc.add_paragraph("None.")
                p.runs[0].italic = True
            else:
                for item in items:
                    doc.add_paragraph(item, style='List Bullet')

    # ---- FOOTER ----
    doc.add_paragraph()
    footer_p = doc.add_paragraph()
    footer_run = footer_p.add_run(
        f"Generated from Fireflies transcript {meeting['transcript_id']} · "
        f"Classified as {meeting_type} "
        f"({meeting.get('type_reasoning', '')})"
    )
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = GREY
    footer_run.italic = True

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _render_todo_section(doc: Document, title: str, todos: list):
    if not todos:
        return
    _add_heading(doc, title, level=2)

    # Sort by owner alphabetically, then by description for stable ordering within owner
    sorted_todos = sorted(
        todos,
        key=lambda t: (t.get('owner', 'zzz').lower(), t.get('description', ''))
    )

    tbl = doc.add_table(rows=1, cols=4)
    tbl.style = 'Light Grid Accent 1'
    tbl.autofit = False
    tbl.allow_autofit = False

    hdr = tbl.rows[0].cells
    hdr[0].text = "Owner"
    hdr[1].text = "Description"
    hdr[2].text = "Due"
    hdr[3].text = "Source"
    for cell in hdr:
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True

    # Column widths that add to ~16cm
    col_widths = [Cm(2.8), Cm(6.5), Cm(2.0), Cm(4.7)]

    # Set header widths
    for i, w in enumerate(col_widths):
        hdr[i].width = w

    last_owner = None
    for todo in sorted_todos:
        row = tbl.add_row()
        owner = todo.get('owner', '')
        # Blank out repeated owner cell to visually group the same owner's items
        row.cells[0].text = '' if owner == last_owner else owner
        row.cells[1].text = todo.get('description', '')
        due = todo.get('due_date')
        row.cells[2].text = due if due else "—"
        row.cells[3].text = f'"{todo.get("source_quote", "")}"'

        # Widths again per row
        for i, w in enumerate(col_widths):
            row.cells[i].width = w

        last_owner = owner